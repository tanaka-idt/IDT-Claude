#!/usr/bin/env python3
"""
Builds the "DCS FY27 Initiatives" spreadsheet: the FY27 counterpart of the
FY26 sheet, with the same three tabs (Tech, Business, Timelines) plus a
Sequence tab with the computed dates and a Capacity tab with the assumptions.

The workbook is written with openpyxl and then uploaded to Drive with
conversion to a native Google Sheet, into the same folder as the FY26 sheet.
The FY26 sheet is never touched.

Usage:
    python build_fy27_sheet.py            # writes the xlsx locally only
    python build_fy27_sheet.py --upload   # also creates the Google Sheet
"""

import argparse
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from fy27_plan_data import (INITIATIVES, CARRY_OVERS, TECH_DEBT_ROWS, TEAM, PHASES,
                            PHASE_COLOR, PHASE_LONG, QUARTERS, QUARTER_LABELS, WEEKS,
                            SIZE_FTE, TECH_RESERVE, SPRINT_WEEKS, ASANA_PROJECT,
                            FY26_SHEET, asana_url, be_fte, app_fte, phase_weeks,
                            week_start, week_label, sprint_of_week, fmt)
from fy27_schedule import schedule, summary_rows

OUT = Path(__file__).parent / "DCS_FY27_Initiatives.xlsx"
FOLDER_ID = "1PM0j7sUc2cG5tVYUC5yojA7E95x-vBDi"      # folder that holds the FY26 sheet
TITLE = "DCS FY27 Initiatives"

HEADER_FILL = PatternFill("solid", fgColor="0B5394")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SECTION_FILL = PatternFill("solid", fgColor="D9D9D9")
SUB_FILL = PatternFill("solid", fgColor="EFEFEF")
GREY = PatternFill("solid", fgColor="B7B7B7")
LIGHT = PatternFill("solid", fgColor="D9D9D9")
RISK_FILL = PatternFill("solid", fgColor="F4CCCC")
THIN = Side(style="thin", color="BFBFBF")
MED = Side(style="medium", color="666666")
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

SCALE_NOTE = ("<<< 1 FTE >>>\nS: 1 sprint\nM: 2 sprints\nL: 4 sprints\n"
              "XL: 9 sprints\nXXL: 15 sprints")
LANES = TEAM["BE"]["lanes"]
SPRINTS = WEEKS // SPRINT_WEEKS


def fill(hexcolor):
    return PatternFill("solid", fgColor=hexcolor)


def phase_label_size(i):
    order = ["-", "S", "M", "L", "XL"]
    return max(i["be"], i["app"], key=order.index)


# ------------------------------------------------------------ Tech / Business --

def write_estimates(ws, business):
    """Tech and Business share a layout; Business adds the calendar-sprint
    columns (FTE sprints divided by the lanes available), as FY26 did."""
    cols = ["Order", "Initiatives", "Description", "New description (revised 17 Sep 2026)", "BE Estimation", "FTE sprints"]
    if business:
        cols += ["BE calendar sprints (2 lanes)"]
    cols += ["APP Estimation", "FTE sprints"]
    if business:
        cols += ["App calendar sprints (2 lanes)"]
    cols += ["BE + APP", "FTE sprints", "Design", "Priority", "Quarter",
             "Cross Dependency", "Comments", "Asana"]
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(1, c)
        cell.fill, cell.font, cell.alignment = HEADER_FILL, HEADER_FONT, CENTER
    ws.row_dimensions[1].height = 34

    # duplicate header names ("FTE sprints"), so indices come from the anchors
    newdesc_c = cols.index("New description (revised 17 Sep 2026)") + 1
    be_size_c = cols.index("BE Estimation") + 1
    be_fte_c = be_size_c + 1
    app_size_c = cols.index("APP Estimation") + 1
    app_fte_c = app_size_c + 1
    tot_size_c = cols.index("BE + APP") + 1
    tot_fte_c = tot_size_c + 1
    design_c = cols.index("Design") + 1
    prio_c, q_c, dep_c, com_c, asana_c = (cols.index("Priority") + 1, cols.index("Quarter") + 1,
                                          cols.index("Cross Dependency") + 1, cols.index("Comments") + 1,
                                          cols.index("Asana") + 1)

    def section(title, row):
        ws.cell(row, 1, title).font = Font(bold=True)
        for c in range(1, len(cols) + 1):
            ws.cell(row, c).fill = SECTION_FILL
        ws.cell(row, be_size_c, SCALE_NOTE).alignment = WRAP
        ws.cell(row, app_size_c, SCALE_NOTE).alignment = WRAP
        ws.cell(row, tot_size_c, "Total estimation\n(BE + APP)").alignment = WRAP
        ws.row_dimensions[row].height = 80

    def item(i, row):
        ws.cell(row, 1, f"{i['quarter']} · #{i['order']}")
        ws.cell(row, 2, i["name"]).font = Font(bold=True)
        ws.cell(row, 3, i["summary"] + "\n\n" + i["description"])
        ws.cell(row, newdesc_c, i["long"])
        ws.cell(row, be_size_c, i["be"] if not i.get("be_left") else f"{i['be']} ({i['be_left']} left)")
        ws.cell(row, be_fte_c, be_fte(i))
        if business:
            ws.cell(row, be_fte_c + 1, f"={get_column_letter(be_fte_c)}{row}/{LANES}")
        ws.cell(row, app_size_c, i["app"] if i["app"] != "-" else "NA")
        ws.cell(row, app_fte_c, app_fte(i))
        if business:
            ws.cell(row, app_fte_c + 1, f"={get_column_letter(app_fte_c)}{row}/{LANES}")
        ws.cell(row, tot_size_c, phase_label_size(i))
        ws.cell(row, tot_fte_c, f"={get_column_letter(be_fte_c)}{row}+{get_column_letter(app_fte_c)}{row}")
        ws.cell(row, design_c, i["design"] if i["design"] != "-" else "NA")
        ws.cell(row, prio_c, i["priority"])
        ws.cell(row, q_c, i["quarter"])
        ws.cell(row, dep_c, i["deps"])
        ws.cell(row, com_c, "PM first-pass estimate; to be validated by engineering.")
        c = ws.cell(row, asana_c, f'=HYPERLINK("{asana_url(i["asana"][0])}","Asana task")')
        c.font = Font(color="1155CC", underline="single")
        for cc in range(1, len(cols) + 1):
            ws.cell(row, cc).alignment = WRAP
        ws.row_dimensions[row].height = 400

    row = 2
    first_item = None
    for title, goal in (("1) Board Goals", "Board Goal"), ("2) Business Goals", "Business Goal")):
        section(title, row)
        row += 1
        for i in sorted([x for x in INITIATIVES if x["goal"] == goal], key=lambda x: x["order"]):
            first_item = first_item or row
            item(i, row)
            row += 1
    section("3) Tech Debt and company tech goals", row)
    row += 1
    for name, desc, be, app in TECH_DEBT_ROWS:
        ws.cell(row, 2, name).font = Font(bold=True)
        ws.cell(row, 3, desc)
        ws.cell(row, be_size_c, be)
        ws.cell(row, be_fte_c, f"={LANES + 1}*{SPRINTS}*{TECH_RESERVE}")
        ws.cell(row, app_size_c, app)
        ws.cell(row, app_fte_c, f"={LANES + 1}*{SPRINTS}*{TECH_RESERVE}")
        ws.cell(row, tot_size_c, "XL")
        ws.cell(row, tot_fte_c, f"={get_column_letter(be_fte_c)}{row}+{get_column_letter(app_fte_c)}{row}")
        ws.cell(row, com_c, "Reserved capacity (35%), the FY26 convention. Not scheduled per item.")
        for cc in range(1, len(cols) + 1):
            ws.cell(row, cc).alignment = WRAP
        ws.row_dimensions[row].height = 60
        tech_row = row
        row += 1

    last = row - 1
    row += 1
    B, A = get_column_letter(be_fte_c), get_column_letter(app_fte_c)
    T = get_column_letter(tot_fte_c)
    lines = [
        ("TOTAL demand (initiatives only)", f"=SUM({B}{first_item}:{B}{tech_row - 2})",
         f"=SUM({A}{first_item}:{A}{tech_row - 2})", f"=SUM({T}{first_item}:{T}{tech_row - 2})"),
        ("Tech reserve (35%)", f"={B}{tech_row}", f"={A}{tech_row}", f"={T}{tech_row}"),
        (f"Raw capacity: 3 BE and 3 App devs x {SPRINTS} sprints", f"=3*{SPRINTS}", f"=3*{SPRINTS}", None),
        ("Capacity left for initiatives (raw minus reserve)", None, None, None),
        ("Over (negative) or under capacity", None, None, None),
    ]
    r0 = row
    for k, (label, b, a, t) in enumerate(lines):
        ws.cell(row, 2, label).font = Font(bold=True)
        if b: ws.cell(row, be_fte_c, b)
        if a: ws.cell(row, app_fte_c, a)
        if t: ws.cell(row, tot_fte_c, t)
        row += 1
    ws.cell(r0 + 3, be_fte_c, f"={B}{r0 + 2}-{B}{r0 + 1}")
    ws.cell(r0 + 3, app_fte_c, f"={A}{r0 + 2}-{A}{r0 + 1}")
    ws.cell(r0 + 4, be_fte_c, f"={B}{r0 + 3}-{B}{r0}")
    ws.cell(r0 + 4, app_fte_c, f"={A}{r0 + 3}-{A}{r0}")
    for c in (be_fte_c, app_fte_c):
        ws.cell(r0 + 4, c).font = Font(bold=True)
    row += 1
    ws.cell(row, 2, "Sizing scale and colours follow the FY26 sheet:").font = Font(italic=True)
    c = ws.cell(row, 3, f'=HYPERLINK("{FY26_SHEET}","DCS FY26 Initiatives")')
    c.font = Font(color="1155CC", underline="single")
    row += 1
    ws.cell(row, 2, "Source board:").font = Font(italic=True)
    c = ws.cell(row, 3, f'=HYPERLINK("{ASANA_PROJECT}","DCS FY27 Asana board")')
    c.font = Font(color="1155CC", underline="single")

    widths = {1: 12, 2: 34, 3: 60, newdesc_c: 90, be_size_c: 14, be_fte_c: 9, app_size_c: 14, app_fte_c: 9,
              tot_size_c: 14, tot_fte_c: 9, design_c: 9, prio_c: 9, q_c: 8,
              dep_c: 34, com_c: 30, asana_c: 12}
    if business:
        widths[be_fte_c + 1] = 12
        widths[app_fte_c + 1] = 12
    for c, w in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "C2"


# ---------------------------------------------------------------- Timelines ---

def write_timelines(ws, sched):
    rows = summary_rows(sched)
    max_week = max(r["end"] for r in rows)
    n_weeks = max(WEEKS, max_week)
    W0 = 18                                       # column R, as in FY26
    last_col = W0 + n_weeks - 1

    # legend, top left so it stays visible with frozen panes
    ws.cell(1, 3, "Legend").font = Font(bold=True)
    legend = [("PM", TEAM["PM"]["label"]), ("Design", TEAM["Design"]["label"]),
              ("TPM", TEAM["TPM"]["label"]), ("BE", TEAM["BE"]["label"]),
              ("App", TEAM["App"]["label"]), ("QA", TEAM["QA"]["label"])]
    for k, (ph, who) in enumerate(legend, start=2):
        ws.cell(k, 2, who)
        c = ws.cell(k, 3, PHASE_LONG[ph])
        c.fill = fill(PHASE_COLOR[ph])
    ws.cell(8, 2, "FY26 carry-over")
    ws.cell(8, 3, "Work in flight when FY27 opens").fill = fill(PHASE_COLOR["CarryOver"])
    ws.cell(9, 2, "Past FY27")
    ws.cell(9, 3, "Lands in FY28 at the 35% tech reserve").fill = RISK_FILL
    ws.cell(1, 5, f"Week 0 = {fmt(week_start(0))}. Sprints are two weeks; S76 is running "
                  f"when the FY opens. PM may start up to 8 weeks before its quarter.").alignment = WRAP
    ws.merge_cells(start_row=1, start_column=5, end_row=3, end_column=16)

    # header rows 10-14: year, quarter, month, sprint, ISO week
    for w in range(n_weeks):
        c = W0 + w
        d = week_start(w) + __import__("datetime").timedelta(days=3)   # mid-week date names the month
        prev = week_start(w - 1) + __import__("datetime").timedelta(days=3) if w else None
        if w == 0 or d.year != prev.year:
            ws.cell(10, c, d.year).font = Font(bold=True)
        if w == 0 or d.month != prev.month:
            ws.cell(12, c, d.strftime("%b")).font = Font(bold=True)
        if w % SPRINT_WEEKS == 0:
            ws.cell(13, c, f"S{sprint_of_week(w)}")
            ws.cell(13, c).alignment = Alignment(horizontal="left")
        ws.cell(14, c, week_label(w)).alignment = Alignment(horizontal="center")
        for r in (10, 11, 12, 13, 14):
            ws.cell(r, c).fill = LIGHT if r in (10, 12, 14) else GREY
        if w >= WEEKS:
            for r in (10, 11, 12, 13, 14):
                ws.cell(r, c).fill = RISK_FILL
        ws.column_dimensions[get_column_letter(c)].width = 2.6
    qs = list(QUARTERS.items()) + [("FY28", WEEKS)]
    for (q, s), (_, e) in zip(qs, qs[1:] + [(None, n_weeks)]):
        if s >= n_weeks:
            continue
        ws.cell(11, W0 + s, q if q != "FY28" else "FY28 (unfunded at 35% reserve)").font = Font(bold=True)
        if e - 1 > s:
            ws.merge_cells(start_row=11, start_column=W0 + s, end_row=11, end_column=W0 + min(e, n_weeks) - 1)
        ws.cell(11, W0 + s).alignment = Alignment(horizontal="center")

    head = ["Goal", "Product", "Initiative", "Priority", "Quarter", "Status", "Size",
            "1-Pager", "PRD", "Data", "Design", "Spike", "BE", "UI", "QA", "Team", "FTE"]
    for k, h in enumerate(head, start=1):
        c = ws.cell(14, k, h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = CENTER
    for k in range(1, 18):
        for r in (10, 11, 12, 13):
            ws.cell(r, k).fill = LIGHT

    def paint(row, ph, start, end, color):
        for w in range(start, end):
            ws.cell(row, W0 + w).fill = fill(color)
        ws.cell(row, W0 + start, ph).font = Font(size=7, color="FFFFFF" if ph in ("BE", "App", "QA") else "000000")

    row = 15
    # carry-overs first
    for c in CARRY_OVERS:
        ph = sched[c["id"]]
        ws.cell(row, 1, "FY26 carry-over").fill = fill(PHASE_COLOR["CarryOver"])
        ws.cell(row, 2, c["product"])
        ws.cell(row, 3, c["name"]).font = Font(bold=True)
        ws.cell(row, 5, "Q1")
        ws.cell(row, 6, "In progress")
        ws.cell(row, 7, "S")
        for k, role in enumerate(["BE", "App", "QA"]):
            r = row + k
            ws.cell(r, 16, role)
            if role in ph:
                s, e, lanes = ph[role]
                ws.cell(r, 17, lanes)
                paint(r, role, s, e, PHASE_COLOR[role])
        for r in range(row, row + 3):
            for k in range(1, 18):
                ws.cell(r, k).border = Border(top=MED if r == row else None)
        row += 3

    for r_ in rows:
        i = r_["i"]
        ph = r_["phases"]
        ws.cell(row, 1, i["goal"])
        ws.cell(row, 2, i["product"])
        c = ws.cell(row, 3, f"#{i['order']} {i['name']}")
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row, 4, i["priority"])
        ws.cell(row, 5, i["quarter"])
        status = "FY28 risk" if r_["spill"] else ("Carry over" if i["fy27"] == "Carry Over" else "Planned")
        sc = ws.cell(row, 6, status)
        if r_["spill"]:
            sc.fill = RISK_FILL
        ws.cell(row, 7, phase_label_size(i))
        for k in range(8, 16):
            ws.cell(row, k, "-" if i["pm_w"] == 0 else "")
        ws.cell(row + 1, 3, f"{fmt(week_start(r_['start']))} to {r_['end_date']}").font = Font(italic=True, size=9)
        ws.cell(row + 2, 3, f"BE {i['be']} ({be_fte(i)} FTE spr) · App {i['app']} ({app_fte(i)} FTE spr) · Design {i['design']}").font = Font(size=9)
        c = ws.cell(row + 3, 3, f'=HYPERLINK("{asana_url(i["asana"][0])}","Asana task")')
        c.font = Font(color="1155CC", underline="single", size=9)
        labels = {"PM": "PM", "Design": "Dsgn", "TPM": "TPM", "BE": "BE", "App": "App", "QA": "QA"}
        for k, role in enumerate(PHASES):
            r = row + k
            ws.cell(r, 16, labels[role])
            if role in ph:
                s, e, lanes = ph[role]
                ws.cell(r, 17, lanes if role in ("BE", "App") else "")
                paint(r, labels[role], s, e, PHASE_COLOR[role])
        for r in range(row, row + 6):
            for k in range(1, 18):
                ws.cell(r, k).border = Border(top=MED if r == row else None)
            for w in range(n_weeks):
                if w % SPRINT_WEEKS == 0:
                    cell = ws.cell(r, W0 + w)
                    cell.border = Border(left=THIN, top=MED if r == row else None)
                if w == WEEKS:
                    cell = ws.cell(r, W0 + w)
                    cell.border = Border(left=MED, top=MED if r == row else None)
        row += 6

    for k, w in {1: 15, 2: 12, 3: 46, 4: 8, 5: 7, 6: 10, 7: 5, 16: 6, 17: 5}.items():
        ws.column_dimensions[get_column_letter(k)].width = w
    for k in range(8, 16):
        ws.column_dimensions[get_column_letter(k)].width = 5
    ws.freeze_panes = "D15"


# ----------------------------------------------------------------- Sequence ---

def write_sequence(ws, sched):
    head = ["Order", "Id", "Initiative", "Goal", "Product", "Priority", "Quarter",
            "BE size", "BE FTE sprints", "App size", "App FTE sprints", "Design",
            "Start", "End", "Status"] + [f"{p} start" for p in PHASES] + [f"{p} end" for p in PHASES] + ["Dependencies", "Asana"]
    ws.append(head)
    for c in range(1, len(head) + 1):
        cell = ws.cell(1, c)
        cell.fill, cell.font, cell.alignment = HEADER_FILL, HEADER_FONT, CENTER
    for r_ in summary_rows(sched):
        i, ph = r_["i"], r_["phases"]
        status = "Lands in FY28 at 35% reserve" if r_["spill"] else "Fits in FY27"
        row = [i["order"], i["id"], i["name"], i["goal"], i["product"], i["priority"],
               i["quarter"], i["be"], be_fte(i), i["app"], app_fte(i), i["design"],
               week_start(r_["start"]), week_start(r_["end"]), status]
        for p in PHASES:
            row.append(week_start(ph[p][0]) if p in ph else "")
        for p in PHASES:
            row.append(week_start(ph[p][1]) if p in ph else "")
        row += [i["deps"], f'=HYPERLINK("{asana_url(i["asana"][0])}","{i["id"]} in Asana")']
        ws.append(row)
        r = ws.max_row
        for c in range(13, 27):
            ws.cell(r, c).number_format = "d mmm yyyy"
        if r_["spill"]:
            ws.cell(r, 15).fill = RISK_FILL
        ws.cell(r, len(head)).font = Font(color="1155CC", underline="single")
    for k, w in {1: 6, 2: 5, 3: 44, 4: 13, 5: 12, 6: 8, 7: 8, 8: 7, 9: 8, 10: 7, 11: 8, 12: 7,
                 13: 12, 14: 12, 15: 24, 28: 50, 29: 14}.items():
        ws.column_dimensions[get_column_letter(k)].width = w
    for k in range(16, 28):
        ws.column_dimensions[get_column_letter(k)].width = 11
    ws.freeze_panes = "D2"


def write_capacity(ws):
    ws.append(["Role", "People", "Parallel lanes used by the plan", "Note"])
    for c in range(1, 5):
        cell = ws.cell(1, c)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for role in PHASES:
        t = TEAM[role]
        ws.append([role, t["people"], t["lanes"], t["note"]])
    ws.append([])
    ws.append(["Assumption", "Value"])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for k, v in [
        ("Fiscal year", f"{fmt(week_start(0))} to {fmt(week_start(WEEKS) - __import__('datetime').timedelta(days=1))}"),
        ("Sprint length", "2 weeks; S76 runs when the FY opens, S101 is the last full FY27 sprint"),
        ("Tech reserve", "35% of BE and App capacity, as in FY26"),
        ("Sizing scale", "S 1, M 2, L 4, XL 9 FTE sprints (FY26 scale)"),
        ("Phase order", "PM, Design, TPM, BE, App, QA in strict sequence per initiative"),
        ("Hand-off rule", "a role that finishes an initiative starts the next one the following week"),
        ("PM lead", "PM may start up to 8 weeks before the initiative's target quarter"),
        ("Estimates", "PM first pass, to be validated by engineering in the Comments column"),
    ]:
        ws.append([k, v])
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 70


def build():
    sched = schedule()
    wb = Workbook()
    ws = wb.active
    ws.title = "Tech"
    write_estimates(ws, business=False)
    write_estimates(wb.create_sheet("Business"), business=True)
    write_timelines(wb.create_sheet("Timelines"), sched)
    write_sequence(wb.create_sheet("Sequence"), sched)
    write_capacity(wb.create_sheet("Capacity"))
    wb.save(OUT)
    print(f"wrote {OUT}")
    return OUT


def upload(path, existing_id=""):
    """Create the Google Sheet, or replace the content of an existing one so its URL survives."""
    from googleapiclient.discovery import build as gbuild
    from googleapiclient.http import MediaFileUpload
    from linkify_refs import get_credentials
    drive = gbuild("drive", "v3", credentials=get_credentials())
    media = MediaFileUpload(str(path),
                            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    if existing_id:
        f = drive.files().update(fileId=existing_id, media_body=media, fields="id,webViewLink").execute()
        print(f"Google Sheet updated in place: {f['webViewLink']}")
        return f
    meta = {"name": TITLE, "mimeType": "application/vnd.google-apps.spreadsheet",
            "parents": [FOLDER_ID]}
    f = drive.files().create(body=meta, media_body=media, fields="id,webViewLink").execute()
    print(f"Google Sheet: {f['webViewLink']}")
    return f


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--upload", action="store_true")
    ap.add_argument("--update", default="", help="spreadsheet id to replace in place (keeps the URL)")
    args = ap.parse_args()
    p = build()
    if args.upload or args.update:
        upload(p, args.update)
